import io
import json
import os
import sqlite3
import sys
import tempfile
import unittest
from contextlib import contextmanager
from datetime import date

from openpyxl import Workbook, load_workbook

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import app as app_module


class EightExperienceContractTest(unittest.TestCase):
    def setUp(self):
        handle = tempfile.NamedTemporaryFile(suffix='.db', delete=False)
        handle.close()
        self.db_path = handle.name
        self.original_get_db = app_module.get_db
        self.original_tokens = dict(app_module._tokens)

        @contextmanager
        def temporary_db():
            db = sqlite3.connect(self.db_path)
            db.row_factory = sqlite3.Row
            try:
                yield db
                db.commit()
            finally:
                db.close()

        app_module.get_db = temporary_db
        app_module._tokens.clear()
        app_module.init_db()
        app_module.migrate_plan_schedules()
        app_module._tokens.update({
            'admin-token': {
                'id': 1, 'role': 'admin', 'roles': ['admin'], 'real_name': '管理员',
            },
            'operator-token': {
                'id': 2, 'role': 'operator', 'roles': ['operator'], 'real_name': '运维甲',
            },
        })
        with temporary_db() as db:
            db.execute('DROP TABLE auth_sessions')
            db.executemany("""INSERT INTO users
                (id,username,password_hash,role,real_name,status)
                VALUES (?,?,?,?,?,'active')""", [
                (1, 'admin-eight', 'x', 'admin', '管理员'),
                (2, 'operator-eight', 'x', 'operator', '运维甲'),
            ])
            db.executemany('INSERT INTO user_roles (user_id,role) VALUES (?,?)',
                           [(1, 'admin'), (2, 'operator')])
            db.execute("""INSERT INTO sites
                (id,code,name,type,status,manager,gps_lat,gps_lng)
                VALUES (1,'SITE-EIGHT','八项验证站','water_quality','normal','运维甲',28.6,115.7)""")
            db.execute('INSERT INTO user_sites (user_id,site_id) VALUES (2,1)')
        self.client = app_module.app.test_client()

    def tearDown(self):
        app_module.get_db = self.original_get_db
        app_module._tokens.clear()
        app_module._tokens.update(self.original_tokens)
        os.unlink(self.db_path)

    @staticmethod
    def headers(token='admin-token'):
        return {'Authorization': f'Bearer {token}'}

    def test_frequency_template_matching_ignores_devices_and_retired_rules(self):
        self.assertEqual(app_module._inspection_frequency('annual'), 'yearly')
        for value in (None, '', 'daily', 'semi_annual', 'high'):
            with self.assertRaises(ValueError, msg=value):
                app_module._inspection_frequency(value)
        with app_module.get_db() as db:
            db.executemany("""INSERT INTO inspection_templates
                (id,template_name,category,frequency,status,sort_order)
                VALUES (?,?,?,?,?,?)""", [
                (501, '周检模板', '水质', 'weekly', 'active', 1),
                (502, '月检模板', '水质', 'monthly', 'active', 2),
                (503, '季检模板', '水质', 'quarterly', 'active', 3),
                (504, '年检模板', '水质', 'yearly', 'active', 4),
                (505, '停用月检模板', '水质', 'monthly', 'inactive', 5),
            ])
            db.executemany("""INSERT INTO inspection_template_items
                (id,template_id,item_name,category,frequency_level,sort_order)
                VALUES (?,?,?,?,?,?)""", [
                (601, 501, '周检项', '水质', '', 1),
                (602, 502, '月检项', '水质', '', 1),
                (603, 503, '季检项', '水质', '', 1),
                (604, 504, '年检项', '水质', '', 1),
                (605, 505, '停用模板项', '水质', '', 1),
            ])
            # Deliberately inactive and device-mismatched legacy rules must not
            # suppress an active frequency template.
            db.execute("""INSERT INTO inspection_configs
                (site_type,device_types,template_id,is_active)
                VALUES ('water_quality','[\"submersible_pump\"]',502,0)""")
            db.executemany("""INSERT INTO device_shadows
                (device_code,device_name,device_type,site_id,status,management_scope)
                VALUES (?,?,?,?,?,?)""", [
                ('PH-1', '不同类型设备', 'ph_meter', 1, 'online', 'managed'),
                ('P-OLD', '退役泵', 'submersible_pump', 1, 'offline', 'retired'),
            ])

        for frequency, item_id in (
                ('weekly', 601), ('monthly', 602),
                ('quarterly', 603), ('yearly', 604), ('annual', 604)):
            matched = self.client.get(
                f'/api/inspection-v2/configs/match?site_id=1&schedule_type={frequency}',
                headers=self.headers('operator-token'))
            self.assertEqual(matched.status_code, 200, matched.json)
            self.assertEqual([item['id'] for item in matched.json['items']], [item_id])
            self.assertEqual(matched.json['device_types'], [])
            self.assertEqual(matched.json['schedule_type'],
                             'yearly' if frequency == 'annual' else frequency)
            self.assertEqual(matched.json['matched_templates'][0]['frequency'],
                             'yearly' if frequency == 'annual' else frequency)

        with app_module.get_db() as db:
            db.execute('DELETE FROM device_shadows')
        no_devices = self.client.get(
            '/api/inspection-v2/configs/match?site_id=1&schedule_type=monthly',
            headers=self.headers('operator-token'))
        self.assertEqual([item['id'] for item in no_devices.json['items']], [602])
        invalid = self.client.get(
            '/api/inspection-v2/configs/match?site_id=1&schedule_type=daily',
            headers=self.headers('operator-token'))
        self.assertEqual((invalid.status_code, invalid.json['code']),
                         (400, 'PLAN_SCHEDULE_TYPE_INVALID'))

    def test_retired_inspection_config_methods_are_read_only(self):
        with app_module.get_db() as db:
            db.execute("DELETE FROM inspection_configs")
            db.execute("""INSERT INTO inspection_templates
                (id,template_name,category,frequency,status,sort_order)
                VALUES (520,'历史规则模板','水质','monthly','active',1)""")
            db.execute("""INSERT INTO inspection_configs
                (id,site_type,device_types,template_id,is_active,remark)
                VALUES (720,'water_quality','[\"legacy_probe\"]',520,1,'保留追溯')""")
            before = [tuple(row) for row in db.execute(
                'SELECT * FROM inspection_configs ORDER BY id').fetchall()]
            schedules_before = db.execute(
                'SELECT COUNT(*) FROM inspection_schedules').fetchone()[0]

        requests = (
            self.client.get('/api/inspection-v2/configs', headers=self.headers()),
            self.client.post('/api/inspection-v2/configs', headers=self.headers(), json={
                'template_id': 520, 'device_types': ['ph_meter'],
            }),
            self.client.put('/api/inspection-v2/configs/720', headers=self.headers(), json={
                'is_active': False,
            }),
            self.client.delete('/api/inspection-v2/configs/720', headers=self.headers()),
        )
        for response in requests:
            self.assertEqual(response.status_code, 410, response.json)
            self.assertEqual(response.json['code'], 'INSPECTION_CONFIG_RULES_RETIRED')
        with app_module.get_db() as db:
            after = [tuple(row) for row in db.execute(
                'SELECT * FROM inspection_configs ORDER BY id').fetchall()]
            schedules_after = db.execute(
                'SELECT COUNT(*) FROM inspection_schedules').fetchone()[0]
        self.assertEqual(after, before)
        self.assertEqual(schedules_after, schedules_before)
        legacy_init = self.client.post(
            '/api/inspection-v2/schedules/init', headers=self.headers())
        self.assertEqual((legacy_init.status_code, legacy_init.json['code']),
                         (410, 'LEGACY_PLAN_RETIRED'))
        with app_module.get_db() as db:
            self.assertEqual(db.execute(
                'SELECT COUNT(*) FROM inspection_schedules').fetchone()[0], schedules_before)
        stats = self.client.get('/api/inspection-v2/stats', headers=self.headers())
        self.assertEqual(stats.status_code, 200, stats.json)
        self.assertEqual(stats.json['total_configs'], 0)

    def test_plan_context_and_selection_ignore_retired_rules_and_keep_history(self):
        with app_module.get_db() as db:
            db.executemany("""INSERT INTO inspection_templates
                (id,template_name,category,frequency,status,sort_order)
                VALUES (?,?,?,?,?,?)""", [
                (522, '有效月检模板', '水质', 'monthly', 'active', 1),
                (523, '停用月检模板', '水质', 'monthly', 'inactive', 2),
                (524, '有效周检模板', '水质', 'weekly', 'active', 3),
            ])
            db.executemany("""INSERT INTO inspection_template_items
                (id,template_id,item_name,category,frequency_level,sort_order)
                VALUES (?,?,?,?,?,?)""", [
                (622, 522, '有效月检项', '水质', '', 1),
                (623, 523, '停用模板项', '水质', '', 1),
                (624, 524, '有效周检项', '水质', '', 1),
            ])
            db.execute("""INSERT INTO inspection_configs
                (site_type,device_types,template_id,is_active,remark)
                VALUES ('water_quality','[\"missing_device\"]',522,0,'旧规则不得过滤')""")
            db.execute("""INSERT INTO plan_schedules
                (id,user_id,schedule_type,period_start,period_end,plan_data,
                 vehicle_days,spare_parts,work_order_ids,status,tasks_generated)
                VALUES (799,2,'monthly','2026-08-20','2026-08-20',?,
                        '{}','[]','[]','draft',0)""",
                       (json.dumps({'2026-08-20': {'sites': [1],
                                                  'inspection_items': {'1': [622]}}}),))
            db.execute("""INSERT INTO insp_plan_items
                (id,plan_id,site_id,template_id,item_name,result)
                VALUES (899,999,1,522,'历史快照名称','normal')""")
            history_before = tuple(db.execute(
                'SELECT template_id,item_name,result FROM insp_plan_items WHERE id=899').fetchone())
            schedule_count = db.execute('SELECT COUNT(*) FROM plan_schedules').fetchone()[0]

        detail = self.client.get('/api/plan-schedules/799', headers=self.headers())
        suggestions = self.client.get(
            '/api/plan-schedules/suggestions?site_ids=1&schedule_type=monthly',
            headers=self.headers())
        self.assertEqual(detail.status_code, 200, detail.json)
        self.assertEqual(suggestions.status_code, 200, suggestions.json)
        expected_context = [{
            'site_id': 1,
            'site_name': '八项验证站',
            'template_name': '有效月检模板',
            'description': '',
            'item_count': 1,
        }]
        self.assertEqual(detail.json['template_context'], expected_context)
        self.assertEqual(suggestions.json['template_context'], expected_context)

        invalid_selections = (
            ('weekly', 622),
            ('monthly', 623),
        )
        for schedule_type, item_id in invalid_selections:
            response = self.client.post('/api/plan-schedules',
                                        headers=self.headers('operator-token'), json={
                'schedule_type': schedule_type,
                'period_start': '2026-08-21',
                'period_end': '2026-08-21',
                'plan_data': {'2026-08-21': {
                    'sites': [1], 'inspection_items': {'1': [item_id]},
                }},
                'vehicle_exception_reason': '本计划无需用车',
            })
            self.assertEqual(response.status_code, 400, response.json)
            self.assertEqual(response.json['code'], 'PLAN_INSPECTION_ITEM_SELECTION_INVALID')
        with app_module.get_db() as db:
            self.assertEqual(db.execute('SELECT COUNT(*) FROM plan_schedules').fetchone()[0],
                             schedule_count)
            history_after = tuple(db.execute(
                'SELECT template_id,item_name,result FROM insp_plan_items WHERE id=899').fetchone())
        self.assertEqual(history_after, history_before)

    def test_period_shrink_rejects_without_silently_dropping_dates(self):
        plan_data = {'2026-08-22': {'sites': [1], 'notes': '保留我'}}
        with app_module.get_db() as db:
            db.execute("""INSERT INTO plan_schedules
                (id,user_id,schedule_type,period_start,period_end,plan_data,vehicle_days,
                 spare_parts,work_order_ids,status,version,remarks,tasks_generated)
                VALUES (701,2,'weekly','2026-08-18','2026-08-24',?,'{}','[]','[]',
                        'draft',3,'原备注',0)""", (json.dumps(plan_data, ensure_ascii=False),))
        response = self.client.put('/api/plan-schedules/701',
                                   headers=self.headers('operator-token'), json={
            'version': 3, 'period_start': '2026-08-18', 'period_end': '2026-08-20',
        })
        self.assertEqual(response.status_code, 409, response.json)
        self.assertEqual(response.json['code'], 'PLAN_DATES_OUTSIDE_PERIOD')
        self.assertEqual(response.json['dates'], ['2026-08-22'])
        with app_module.get_db() as db:
            row = db.execute('SELECT period_end,plan_data,version FROM plan_schedules WHERE id=701').fetchone()
        self.assertEqual((row['period_end'], row['version']), ('2026-08-24', 3))
        self.assertEqual(json.loads(row['plan_data']), plan_data)

    def test_create_and_validate_share_strict_date_and_owner_contracts(self):
        base = {
            'schedule_type': 'weekly',
            'period_start': '2026-08-18',
            'period_end': '2026-08-24',
            'plan_data': {},
            'vehicle_exception_reason': '本计划无需用车',
        }
        with app_module.get_db() as db:
            initial_count = db.execute('SELECT COUNT(*) FROM plan_schedules').fetchone()[0]
        invalid_payloads = [
            {**base, 'period_start': '2026-13-01'},
            {**base, 'period_start': '2026-08-25', 'period_end': '2026-08-24'},
            {**base, 'plan_data': {'not-a-date': {'notes': '真实安排'}}},
        ]
        for payload in invalid_payloads:
            created = self.client.post('/api/plan-schedules', headers=self.headers('operator-token'), json=payload)
            validated = self.client.post('/api/plan-schedules/validate', headers=self.headers('operator-token'), json=payload)
            self.assertEqual(created.status_code, 400, created.json)
            self.assertEqual(validated.status_code, 400, validated.json)
        with app_module.get_db() as db:
            self.assertEqual(db.execute('SELECT COUNT(*) FROM plan_schedules').fetchone()[0], initial_count)

        forbidden = self.client.post('/api/plan-schedules/validate', headers=self.headers('operator-token'),
                                     json={**base, 'user_id': 1})
        self.assertEqual(forbidden.status_code, 403, forbidden.json)
        admin_other = self.client.post('/api/plan-schedules/validate', headers=self.headers(),
                                       json={**base, 'user_id': 2})
        self.assertEqual(admin_other.status_code, 200, admin_other.json)

        custom_date = self.client.post('/api/plan-schedules', headers=self.headers('operator-token'), json={
            **base,
            'plan_data': {'2026-08-20': {'sites': [1], 'notes': '非整周自定义执行日'}},
        })
        self.assertEqual(custom_date.status_code, 201, custom_date.json)
        self.assertEqual(custom_date.json['period_start'], '2026-08-18')

    def test_legacy_frequency_values_are_flagged_instead_of_reinterpreted(self):
        for value in ('high', 'mid', 'low', 'daily', 'semi_annual'):
            self.assertEqual(app_module._mobile_frequency_display(value), ('需管理员维护', True))
        self.assertEqual(app_module._mobile_frequency_display('weekly'), ('每周', False))
        self.assertEqual(app_module._mobile_frequency_display('yearly'), ('每年', False))

    def test_due_suggestion_is_retired_without_writes_or_history_loss(self):
        today = date.today().isoformat()
        with app_module.get_db() as db:
            db.execute("""INSERT INTO inspection_templates
                (id,template_name,category,frequency,status,sort_order)
                VALUES (510,'到期月检','水质','monthly','active',1)""")
            db.execute("""INSERT INTO inspection_template_items
                (id,template_id,item_name,category,frequency_level,sort_order)
                VALUES (610,510,'水质到期项','水质','',1)""")
            db.execute("""INSERT INTO inspection_configs
                (site_type,device_types,template_id,is_active)
                VALUES ('water_quality','[]',510,1)""")
            db.execute("""INSERT INTO inspection_schedules
                (id,site_id,template_id,template_item_id,frequency,next_due_date,status)
                VALUES (710,1,510,610,'monthly',?,'active')""", (today,))
            db.execute('INSERT INTO user_roles (user_id,role) VALUES (1,\'operator\')')
            db.execute('INSERT INTO user_sites (user_id,site_id) VALUES (1,1)')
            app_module._create_notification(
                2, 'inspection_due_suggestion', 710, '历史到期巡检建议',
                '该通知创建于功能停用前', db=db)
            before = {
                'notifications': db.execute('SELECT COUNT(*) FROM notifications').fetchone()[0],
                'schedules': db.execute('SELECT COUNT(*) FROM plan_schedules').fetchone()[0],
                'events': db.execute('SELECT COUNT(*) FROM plan_schedule_events').fetchone()[0],
                'execution_plans': db.execute('SELECT COUNT(*) FROM insp_plans').fetchone()[0],
            }

        self.assertEqual(app_module.create_due_schedule_drafts_job(), {'disabled': True})
        with app_module.get_db() as db:
            self.assertEqual({
                'notifications': db.execute('SELECT COUNT(*) FROM notifications').fetchone()[0],
                'schedules': db.execute('SELECT COUNT(*) FROM plan_schedules').fetchone()[0],
                'events': db.execute('SELECT COUNT(*) FROM plan_schedule_events').fetchone()[0],
                'execution_plans': db.execute('SELECT COUNT(*) FROM insp_plans').fetchone()[0],
            }, before)
            historical = db.execute("""SELECT is_read FROM notifications
                WHERE user_id=2 AND source_type='inspection_due_suggestion' AND source_id=710""").fetchone()
            self.assertIsNotNone(historical)
            self.assertEqual(historical['is_read'], 0)

        for method in ('get', 'post'):
            kwargs = {'json': {'action': 'create_draft', 'user_id': 2}} if method == 'post' else {}
            response = getattr(self.client, method)(
                '/api/plan-schedules/draft-recommendations',
                headers=self.headers('operator-token'), **kwargs)
            self.assertEqual(response.status_code, 410, response.json)
            self.assertEqual(response.json['code'], 'INSPECTION_DUE_SUGGESTION_RETIRED')
        legacy_generator = self.client.post(
            '/api/inspection-v2/plans/generate', headers=self.headers('operator-token'))
        self.assertEqual(legacy_generator.status_code, 410, legacy_generator.json)
        self.assertNotIn('到期巡检建议', legacy_generator.json['error'])
        with app_module.get_db() as db:
            self.assertEqual({
                'notifications': db.execute('SELECT COUNT(*) FROM notifications').fetchone()[0],
                'schedules': db.execute('SELECT COUNT(*) FROM plan_schedules').fetchone()[0],
                'events': db.execute('SELECT COUNT(*) FROM plan_schedule_events').fetchone()[0],
                'execution_plans': db.execute('SELECT COUNT(*) FROM insp_plans').fetchone()[0],
            }, before)

    def test_generated_tasks_are_counted_by_date_and_site_with_user_status(self):
        with app_module.get_db() as db:
            db.execute("""INSERT INTO plan_schedules
                (id,user_id,schedule_type,period_start,period_end,plan_data,vehicle_days,
                 spare_parts,work_order_ids,status,tasks_generated)
                VALUES (702,2,'weekly','2026-08-18','2026-08-24','{}','{}','[]','[]',
                        'approved',1)""")
            db.executemany("""INSERT INTO insp_plans
                (id,plan_name,assignee,assignee_id,period,generate_date,status,plan_schedule_id)
                VALUES (?,?,?,?,?,?,?,702)""", [
                (801, '第一日', '运维甲', 2, 'weekly', '2026-08-19', 'active'),
                (802, '第二日', '运维甲', 2, 'weekly', '2026-08-20', 'active'),
            ])
            db.executemany("""INSERT INTO insp_plan_items
                (id,plan_id,site_id,item_name,result,execution_status)
                VALUES (?,?,?,?,?,'active')""", [
                (901, 801, 1, '检查A', 'normal'),
                (902, 801, 1, '检查B', None),
                (903, 802, 1, '检查C', 'normal'),
            ])
            rows = app_module._ps_execution_site_rows(db, 702)
            change_rows = app_module._ps_execution_site_rows(db, 702, 'change_submitted')
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0]['status'], 'partial')
        self.assertEqual(rows[1]['status'], 'completed')
        self.assertEqual([row['status'] for row in change_rows], ['change_pending', 'change_pending'])
        self.assertEqual(app_module._ps_schedule_execution_status('approved', 'active', rows), {
            'execution_status': 'partial', 'execution_status_cn': '部分完成',
        })
        self.assertEqual(app_module._ps_schedule_execution_status(
            'approved', 'active', [{'status': 'pending'}]), {
                'execution_status': 'pending', 'execution_status_cn': '待执行',
            })
        self.assertEqual(app_module._ps_schedule_execution_status(
            'approved', 'active', [{'status': 'completed'}, {'status': 'completed'}]), {
                'execution_status': 'completed', 'execution_status_cn': '已完成',
            })
        self.assertEqual(app_module._ps_schedule_execution_status(
            'change_submitted', 'active', rows), {
                'execution_status': 'change_pending', 'execution_status_cn': '变更待审',
            })
        self.assertEqual(app_module._ps_schedule_execution_status('approved', 'active', []), {
            'execution_status': 'pending', 'execution_status_cn': '待执行',
        })
        self.assertEqual(app_module._ps_schedule_execution_status('approved', 'rework', []), {
            'execution_status': 'rework', 'execution_status_cn': '需整改',
        })

        detail = self.client.get('/api/plan-schedules/702', headers=self.headers())
        listed = self.client.get('/api/plan-schedules', headers=self.headers())
        self.assertEqual(detail.status_code, 200, detail.json)
        self.assertEqual((detail.json['execution_status'], detail.json['execution_status_cn']),
                         ('partial', '部分完成'))
        listed_row = next(item for item in listed.json if item['id'] == 702)
        self.assertEqual((listed_row['execution_status'], listed_row['execution_status_cn']),
                         ('partial', '部分完成'))

        with app_module.get_db() as db:
            db.execute("UPDATE plan_schedules SET status='change_submitted' WHERE id=702")
        changed = self.client.get('/api/plan-schedules/702', headers=self.headers())
        self.assertEqual((changed.json['execution_status'], changed.json['execution_status_cn']),
                         ('change_pending', '变更待审'))

    def test_plan_list_batches_execution_status_without_per_schedule_detail_queries(self):
        schedules = [
            (710, 'approved', 'active'),
            (711, 'approved', 'active'),
            (712, 'approved', 'active'),
            (713, 'change_submitted', 'active'),
            (714, 'approved', 'rework'),
        ]
        with app_module.get_db() as db:
            db.executemany("""INSERT INTO plan_schedules
                (id,user_id,schedule_type,period_start,period_end,plan_data,vehicle_days,
                 spare_parts,work_order_ids,status,tasks_generated,field_status)
                VALUES (?,2,'weekly','2026-08-18','2026-08-24','{}','{}','[]','[]',?,1,?)""",
                           schedules)
            db.executemany("""INSERT INTO insp_plans
                (id,plan_name,assignee,assignee_id,period,generate_date,status,plan_schedule_id)
                VALUES (?,?, '运维甲',2,'weekly','2026-08-19',?,?)""", [
                (810, '待执行', 'active', 710),
                (811, '部分完成', 'active', 711),
                (812, '已完成', 'completed', 712),
                (813, '变更待审', 'active', 713),
                (814, '需整改', 'active', 714),
            ])
            db.executemany("""INSERT INTO insp_plan_items
                (plan_id,site_id,item_name,result,execution_status,check_out_time)
                VALUES (?,?,?,?,'active',?)""", [
                (810, 1, '待执行项', None, None),
                (811, 1, '已完成项', 'normal', None),
                (811, 1, '未完成项', None, None),
                (812, 1, '全部完成项', 'normal', '2026-08-19 10:00:00'),
                (813, 1, '变更项', None, None),
                (814, 1, '整改项', None, None),
            ])

        original_batch = app_module._ps_execution_item_summaries
        original_detail = app_module._ps_execution_site_rows
        batch_calls = []

        def counted_batch(db, schedule_ids):
            batch_calls.append(tuple(schedule_ids))
            return original_batch(db, schedule_ids)

        def forbidden_detail(*_args, **_kwargs):
            raise AssertionError('列表不得逐计划调用日期×站点明细聚合')

        app_module._ps_execution_item_summaries = counted_batch
        app_module._ps_execution_site_rows = forbidden_detail
        try:
            response = self.client.get('/api/plan-schedules', headers=self.headers())
            self.assertEqual(response.status_code, 200, response.json)
            self.assertEqual(len(batch_calls), 1)
            by_id = {item['id']: item for item in response.json}
            self.assertEqual({sid: by_id[sid]['execution_status'] for sid, _, _ in schedules}, {
                710: 'pending', 711: 'partial', 712: 'completed',
                713: 'change_pending', 714: 'rework',
            })

            with app_module.get_db() as db:
                db.execute('DELETE FROM plan_schedules')
            batch_calls.clear()
            empty = self.client.get('/api/plan-schedules', headers=self.headers())
            self.assertEqual(empty.status_code, 200, empty.json)
            self.assertEqual(empty.json, [])
            self.assertEqual(batch_calls, [])
        finally:
            app_module._ps_execution_item_summaries = original_batch
            app_module._ps_execution_site_rows = original_detail

    def test_chinese_device_template_round_trips_and_site_archive_reopens(self):
        template = self.client.get('/api/import-templates/devices', headers=self.headers())
        self.assertEqual(template.status_code, 200)
        template_book = load_workbook(io.BytesIO(template.data))
        self.assertEqual(template_book.sheetnames,
                         ['设备导入', '站点参考', '设备类型参考', '状态参考'])
        self.assertEqual(template_book['设备导入']['D1'].value, '设备类型')
        self.assertTrue(template_book['设备导入'].data_validations.count >= 2)
        template_book.close()

        upload = Workbook()
        sheet = upload.active
        sheet.title = '设备导入'
        sheet.append(['站点编码', '设备编码', '设备名称', '设备类型', '设备型号', '厂商', '安装日期', '状态', '备注'])
        sheet.append(['SITE-EIGHT', 'DEV-EIGHT', '现场pH计', 'pH 计', '', '', '2026-08-18', '正常', ''])
        payload = io.BytesIO()
        upload.save(payload)
        payload.seek(0)
        preview = self.client.post('/api/imports/devices/validate', headers=self.headers(),
                                   data={'file': (payload, 'devices.xlsx')},
                                   content_type='multipart/form-data')
        self.assertEqual(preview.status_code, 200, preview.json)
        self.assertEqual(preview.json['valid_count'], 1)
        row = preview.json['rows'][0]
        self.assertEqual((row['data']['device_type'], row['data']['status']), ('ph_meter', 'normal'))
        committed = self.client.post('/api/imports/devices/commit', headers=self.headers(),
                                     json={'rows': preview.json['rows']})
        self.assertEqual(committed.status_code, 200, committed.json)

        with app_module.get_db() as db:
            db.execute("UPDATE sites SET code='001', type='water_quality', status='offline' WHERE id=1")
            db.executemany("""INSERT INTO device_shadows
                (site_id,device_code,device_name,device_type,status) VALUES (1,?,?,?,?)""", [
                ('0001', '高锰酸盐指数在线分析设备', 'codmn_analyzer', 'online'),
                ('0002', '氨氮在线分析设备', 'ammonia_analyzer', 'offline'),
                ('0003', '历史未知设备', 'legacy_probe', 'mystery_status'),
                ('0004', '维护中的pH设备', 'ph_meter', 'maintenance'),
                ('0005', '制度牌与门锁', 'station_facility', 'online'),
                ('0006', 'COD在线分析设备', 'cod_analyzer', 'online'),
            ])

        archive = self.client.get('/api/sites/1/archive/export', headers=self.headers())
        self.assertEqual(archive.status_code, 200)
        archive_book = load_workbook(io.BytesIO(archive.data))
        self.assertEqual(archive_book.sheetnames,
                         ['站点概况', '设备', '巡检工单摘要', '风险信息'])
        device_sheet = archive_book['设备']
        overview = archive_book['站点概况']
        overview_values = {overview.cell(row, 1).value: overview.cell(row, 2).value
                           for row in range(2, overview.max_row + 1)}
        self.assertEqual(overview_values['站点编码'], '001')
        self.assertEqual(overview_values['站点类型'], '水质监测站')
        self.assertEqual(overview_values['状态'], '离线')
        self.assertEqual(overview['B2'].number_format, '@')
        headers = [cell.value for cell in device_sheet[1]]
        code_column = headers.index('设备编码') + 1
        self.assertEqual(device_sheet.cell(2, code_column).value, 'DEV-EIGHT')
        device_rows = {device_sheet.cell(row, code_column).value: row
                       for row in range(2, device_sheet.max_row + 1)}
        self.assertEqual(device_sheet.cell(device_rows['0001'], code_column).number_format, '@')
        self.assertEqual(device_sheet.cell(device_rows['0001'], 3).value, '高锰酸盐分析仪')
        self.assertEqual(device_sheet.cell(device_rows['0001'], 6).value, '在线')
        self.assertEqual(device_sheet.cell(device_rows['0002'], 3).value, '氨氮分析仪')
        self.assertEqual(device_sheet.cell(device_rows['0002'], 6).value, '离线')
        self.assertEqual(device_sheet.cell(device_rows['0003'], 3).value, 'legacy_probe（需维护）')
        self.assertEqual(device_sheet.cell(device_rows['0003'], 6).value, 'mystery_status（需维护）')
        self.assertEqual(device_sheet.cell(device_rows['0004'], 6).value, '维护中')
        self.assertEqual(device_sheet.cell(device_rows['0005'], 3).value, '站房设施')
        self.assertEqual(device_sheet.cell(device_rows['0006'], 3).value, 'COD分析仪')
        self.assertGreaterEqual(device_sheet.column_dimensions['B'].width, 28)
        self.assertGreaterEqual(device_sheet.column_dimensions['C'].width, 24)
        self.assertTrue(device_sheet['B2'].alignment.wrap_text)
        archive_book.close()


if __name__ == '__main__':
    unittest.main()
